"""Builds the bulk-upload Excel template in memory, tailored to the selected
mode_of_operation. No files are written to disk — this avoids the whole class
of bugs where a stale/forgotten regenerated file gets served by mistake.
"""

from __future__ import annotations

from io import BytesIO

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.comments import Comment

# ---------------------------------------------------------------------------
# Column sets per mode
# ---------------------------------------------------------------------------
COMMON_COLUMNS = [
    "part_code", "part_name", "category_name", "image", "part_weight",
]

FIXED_COLUMN_HELP = {
    "part_code": "REQUIRED. Unique identifier for this part (e.g. BOLT-001).\n"
                 "Duplicate part_code rows are skipped on upload.",
    "part_name": "Display name of the part. Falls back to part_code if left blank.",
    "category_name": "Groups this part under a category. Auto-created if new.\n"
                      "Grouping/filtering only — does not select or load any model.",
    "image": "Optional. Leave blank unless your process embeds an image reference here.",
    "part_weight": "Weight of the part, in grams. A single fixed number — no min/max, no instances.",
    "part_co_planarity": "true/false — whether this geometric check applies.",
    "part_parallelity": "true/false — whether this geometric check applies.",
    "part_concentricity": "true/false — whether this geometric check applies.",
}

# Counting: nominal-value-only columns, no min/max, no geometric flags
COUNTING_FAMILIES = ["length", "width", "height", "id", "od"]

# Measurement: all 8 families with full value/min/max
MEASUREMENT_FAMILIES = ["length", "width", "height", "id", "od", "angle", "arch", "sector"]

FAMILY_LABELS = {
    "length": "Length", "width": "Width", "height": "Height",
    "id": "Inner Diameter", "od": "Outer Diameter",
    "angle": "Angle", "arch": "Arch Length", "sector": "Sector",
}

MODE_SLUGS = {
    "Counting": "Counting",
    "Defect Detection": "DefectDetection",
    "Measurement": "Measurement",
}


def _style_header(ws, headers: list[str]):
    hf = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="B71C1C")
    for c in ws[1]:
        c.font = hf
        c.fill = fill
    for i, col in enumerate(headers, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = max(14, len(col) + 2)


def _add_comments(ws, headers: list[str]):
    name_to_col = {name: i for i, name in enumerate(headers, 1)}
    for name, help_text in FIXED_COLUMN_HELP.items():
        if name in name_to_col:
            ws.cell(row=1, column=name_to_col[name]).comment = Comment(help_text, "Template Guide")


def _add_instructions(wb, mode: str, extra_lines: list[str]):
    ws_info = wb.create_sheet("Instructions", 0)
    ws_info.column_dimensions["A"].width = 24
    ws_info.column_dimensions["B"].width = 78

    ws_info.append([f"Bulk Upload Template — {mode}"])
    ws_info["A1"].font = Font(bold=True, size=14)
    ws_info.append([])
    ws_info.append([f"This template is for parts in '{mode}' mode only.",
                     "Every part uploaded from this file will be created with that mode."])
    ws_info.append([])

    ws_info.append(["Fixed Columns"])
    ws_info.cell(row=ws_info.max_row, column=1).font = Font(bold=True, size=12)
    ws_info.append(["Column", "What it means"])
    for c in ws_info[ws_info.max_row]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="EFEFEF")
    fixed_reference = [
        ("part_code", "REQUIRED. Unique ID for the part. Duplicates are skipped."),
        ("part_name", "Display name. Defaults to part_code if left blank."),
        ("category_name", "Groups the part under a category. Auto-created if new."),
        ("image", "Optional. Leave blank unless your process embeds an image reference."),
        ("part_weight", "Weight in grams. One fixed number — no min/max, no instances."),
    ]
    for col, desc in fixed_reference:
        ws_info.append([col, desc])
    ws_info.append([])

    for line in extra_lines:
        ws_info.append([None, line] if isinstance(line, str) else line)


def _nominal_only_headers(families: list[str]) -> list[str]:
    """Just the nominal-value column per family — no _min/_max."""
    return list(families)


def _dimension_headers_with_tolerance(families: list[str]) -> list[str]:
    headers = []
    for base in families:
        headers += [f"{base}1", f"{base}1_min", f"{base}1_max"]
    return headers


def _build_counting(wb, ws) -> list[str]:
    """Counting: mandatory columns + plain nominal-value dimensions only.
    No min/max, no geometric check flags."""
    dim_headers = [f"{base}1" for base in COUNTING_FAMILIES]
    headers = list(COMMON_COLUMNS) + dim_headers
    ws.append(headers)
    ws.append([
        "BOLT-001", "M8 Bolt", "Fasteners", "", 12.5,
        40,   # length1
        8,    # width1
        "",   # height1
        6.2,  # id1
        8.0,  # od1
    ])
    _add_instructions(wb, "Counting", [
        "Dimension Parameters (nominal values only — no min/max for Counting)",
        "  length1, width1, height1, id1, od1 — the nominal size of the part.",
        "'id' = inner diameter, 'od' = outer diameter.",
        "Leave any column blank if it doesn't apply to this part.",
    ])
    return headers


def _build_defect_detection(wb, ws) -> list[str]:
    """Defect Detection: mandatory columns + defect names only.
    No dimension columns, no geometric check flags."""
    headers = list(COMMON_COLUMNS) + ["d1_name", "d2_name"]
    ws.append(headers)
    ws.append([
        "BRACKET-001", "L-Bracket", "Brackets", "", 45.0,
        "dent", "scratch",
    ])
    _add_instructions(wb, "Defect Detection", [
        "Defects",
        "d1_name is one defect. Add more the same way: d2_name, d3_name, etc.",
        "There is no limit to the number of defects per part.",
    ])
    return headers


def _build_measurement(wb, ws) -> list[str]:
    """Measurement: mandatory columns + full value/min/max for all 8 families."""
    headers = list(COMMON_COLUMNS) + _dimension_headers_with_tolerance(MEASUREMENT_FAMILIES)
    ws.append(headers)
    example = {
        "part_code": "GEAR-001", "part_name": "Spur Gear 20T", "category_name": "Gears",
        "part_weight": 85.0,
        "od1": 25.0, "od1_min": 24.9, "od1_max": 25.1,
        "id1": 6.0, "id1_min": 5.95, "id1_max": 6.05,
        "arch1": 3.9, "sector1": 18,
    }
    ws.append([example.get(h, "") for h in headers])
    _add_instructions(wb, "Measurement", [
        "Dimension Parameters (length, width, height, id, od, angle, arch, sector)",
        "Each family supports UNLIMITED numbered instances (length1, length2, ...).",
        "  length1 = nominal value, length1_min / length1_max = tolerance (optional).",
        "Add more by adding columns yourself: length2, length2_min, length2_max, etc.",
        "'id' = inner diameter, 'od' = outer diameter, 'arch' = arch length.",
    ])
    return headers


_BUILDERS = {
    "Counting": _build_counting,
    "Defect Detection": _build_defect_detection,
    "Measurement": _build_measurement,
}


def build_template(mode: str) -> BytesIO:
    """Returns an in-memory .xlsx file tailored to the given mode."""
    if mode not in _BUILDERS:
        raise ValueError(f"Unknown mode: {mode}")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Parts"

    headers = _BUILDERS[mode](wb, ws)
    _style_header(ws, headers)
    _add_comments(ws, headers)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf