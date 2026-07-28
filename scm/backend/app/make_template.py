import os
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

FAMILIES = ["length", "width", "height", "id", "od", "angle", "arch", "sector"]

# ---------------------------------------------------------------------------
# Column headers — NOTE: no mode_of_operation column here anymore.
# Mode is set once, globally, in cell B1 above the table.
# ---------------------------------------------------------------------------
headers = ["part_code", "part_name", "category_name", "image", "part_weight"]
for base in FAMILIES:
    headers += [f"{base}1", f"{base}1_min", f"{base}1_max"]
headers += ["d1_name", "d2_name"]
headers += ["part_co_planarity", "part_parallelity", "part_concentricity"]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Parts"

# ---- Row 1: global mode setting ----
ws["A1"] = "Mode of Operation for this sheet:"
ws["A1"].font = Font(bold=True)
ws["B1"] = "Counting"  # default value — user changes via dropdown
ws["B1"].font = Font(bold=True, color="B71C1C")
ws["B1"].fill = PatternFill("solid", fgColor="FFF3E0")

dv = DataValidation(type="list", formula1='"Counting,Defect Detection,Measurement"', allow_blank=False)
ws.add_data_validation(dv)
dv.add("B1")

ws["B1"].comment = Comment(
    "Choose ONE mode for this entire sheet: Counting, Defect Detection, or "
    "Measurement. Every part row below is uploaded under this mode — there is "
    "no per-row mode column.",
    "Template Guide",
)

# ---- Row 2: blank spacer ----
# (left empty on purpose)

# ---- Row 3: column headers ----
HEADER_ROW = 3
for col_idx, name in enumerate(headers, start=1):
    ws.cell(row=HEADER_ROW, column=col_idx, value=name)

hf = Font(bold=True, color="FFFFFF")
fill = PatternFill("solid", fgColor="B71C1C")
for c in ws[HEADER_ROW]:
    c.font = hf
    c.fill = fill
for i, col in enumerate(headers, 1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = max(14, len(col) + 2)

# ---- Row 4+: example data rows ----
rows = [
    {
        "part_code": "BOLT-001", "part_name": "M8 Bolt", "category_name": "Fasteners",
        "part_weight": 12.5,
        "length1": 40, "length1_min": 39.5, "length1_max": 40.5,
        "width1": 8,
        "od1": 8.0, "od1_min": 7.8, "od1_max": 8.2,
        "part_co_planarity": "false", "part_parallelity": "false", "part_concentricity": "false",
    },
    {
        "part_code": "WASHER-002", "part_name": "M8 Washer", "category_name": "Fasteners",
        "part_weight": 2.1,
        "id1": 8.4, "id1_min": 8.2, "id1_max": 8.6,
        "od1": 16.0, "od1_min": 15.8, "od1_max": 16.2,
        "height1": 1.5,
        "part_co_planarity": "false", "part_parallelity": "false", "part_concentricity": "false",
    },
]
for r_idx, row in enumerate(rows, start=HEADER_ROW + 1):
    for col_idx, name in enumerate(headers, start=1):
        ws.cell(row=r_idx, column=col_idx, value=row.get(name, ""))

# ---------------------------------------------------------------------------
# Header cell comments — hover tooltips for fixed columns
# ---------------------------------------------------------------------------
FIXED_COLUMN_HELP = {
    "part_code": "REQUIRED. Unique identifier for this part (e.g. BOLT-001).\n"
                 "Duplicate part_code rows are skipped on upload.",
    "part_name": "Display name of the part. Falls back to part_code if left blank.",
    "category_name": "Groups this part under a category. Auto-created if new.\n"
                      "Grouping/filtering only — does not select or load any model.",
    "image": "Optional. Leave blank unless your process embeds an image reference here.",
    "part_weight": "Weight of the part, in grams. One fixed number — no min/max, no instances.",
    "part_co_planarity": "true/false — whether this geometric check applies.",
    "part_parallelity": "true/false — whether this geometric check applies.",
    "part_concentricity": "true/false — whether this geometric check applies.",
}
name_to_col = {name: i for i, name in enumerate(headers, 1)}
for name, help_text in FIXED_COLUMN_HELP.items():
    if name in name_to_col:
        ws.cell(row=HEADER_ROW, column=name_to_col[name]).comment = Comment(help_text, "Template Guide")

# ---------------------------------------------------------------------------
# Instructions sheet
# ---------------------------------------------------------------------------
ws_info = wb.create_sheet("Instructions", 0)
ws_info.column_dimensions["A"].width = 24
ws_info.column_dimensions["B"].width = 78

ws_info.append(["How to fill this template"])
ws_info["A1"].font = Font(bold=True, size=14)
ws_info.append([])

ws_info.append(["Step 1 — Set the mode ONCE for the whole sheet"])
ws_info.cell(row=ws_info.max_row, column=1).font = Font(bold=True, size=12)
ws_info.append([None, "On the 'Parts' sheet, cell B1 sets the mode for every part in this file:"])
ws_info.append([None, "  Counting, Defect Detection, or Measurement (pick from the dropdown)."])
ws_info.append([None, "There is no per-row mode column — every row is uploaded under this one mode."])
ws_info.append([None, "To upload parts for a different mode, use a separate file with a different B1 value."])
ws_info.append([])

ws_info.append(["Fixed Columns"])
ws_info.cell(row=ws_info.max_row, column=1).font = Font(bold=True, size=12)
ws_info.append(["Column", "What it means"])
for c in ws_info[ws_info.max_row]:
    c.font = Font(bold=True)
    c.fill = PatternFill("solid", fgColor="EFEFEF")
fixed_reference = [
    ("part_code", "REQUIRED. Unique ID for the part. Duplicates are skipped."),
    ("part_name", "Display name. Defaults to part_code if left blank."),
    ("category_name", "Groups the part under a category. Auto-created if new."),
    ("image", "Optional. Leave blank unless your process embeds an image reference."),
    ("part_weight", "Weight in grams. One fixed number — no min/max, no instances."),
    ("part_co_planarity", "true/false."),
    ("part_parallelity", "true/false."),
    ("part_concentricity", "true/false."),
]
for col, desc in fixed_reference:
    ws_info.append([col, desc])
ws_info.append([])

ws_info.append(["Dimension Parameters (length, width, height, id, od, angle, arch, sector)"])
ws_info.cell(row=ws_info.max_row, column=1).font = Font(bold=True, size=12)
for line in [
    "Each family supports UNLIMITED numbered instances (length1, length2, ...).",
    "  length1 = nominal value, length1_min / length1_max = tolerance (optional).",
    "Add more by adding columns yourself: length2, length2_min, length2_max, etc.",
    "'id' = inner diameter, 'od' = outer diameter, 'arch' = arch length.",
]:
    ws_info.append([None, line])
ws_info.append([])

ws_info.append(["Defects"])
ws_info.cell(row=ws_info.max_row, column=1).font = Font(bold=True, size=12)
for line in [
    "d1_name is one defect. Add more the same way: d2_name, d3_name, etc.",
]:
    ws_info.append([None, line])

base_dir = os.path.dirname(os.path.abspath(__file__))
assets_dir = os.path.join(base_dir, "assets")
os.makedirs(assets_dir, exist_ok=True)
path = os.path.join(assets_dir, "TemplateBulkUpload.xlsx")
wb.save(path)
print("Wrote", path)