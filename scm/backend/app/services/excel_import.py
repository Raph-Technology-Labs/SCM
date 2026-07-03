"""
Excel import: dump an .xlsx of parts into the database and link each part to
its AI model by name.

Expected columns (header row, case-insensitive; extras are ignored):
    part_code            (required)
    part_name            (required)
    model_name           -> linked to ai_models (e.g. 'bolt' => bolt.pt)
    category_id
    mode_of_operation
    part_weight, part_height, part_width
    part_inner_diameter, part_outer_diameter
    part_length, part_angle, part_arch_length
    part_co_planarity, part_parallelity, part_concentricity   (0/1 or true/false)
    parts_metadata

Existing part_codes are skipped (reported in errors) so re-running is safe.
"""

from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.inference import resolve_model_filename
from app.models.db import AIModel, Part
from app.schemas import ImportResult


def _to_bool(v: Any) -> bool:
    if isinstance(v, bool):
        return v
    return str(v).strip().lower() in {"1", "true", "yes", "y", "on"}


def _to_float(v: Any):
    if v is None or str(v).strip() == "":
        return None
    return float(v)


def _to_int(v: Any):
    if v is None or str(v).strip() == "":
        return None
    return int(float(v))


def _to_str(v: Any):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def import_parts_from_xlsx(db: Session, file_bytes: bytes) -> ImportResult:
    wb = load_workbook(BytesIO(file_bytes), data_only=True, read_only=True)
    ws = wb.active

    rows = ws.iter_rows(values_only=True)
    try:
        header = [str(h).strip().lower() if h is not None else "" for h in next(rows)]
    except StopIteration:
        return ImportResult(created_parts=0, created_models=0, linked_parts=0,
                            errors=["Empty spreadsheet"])

    result = ImportResult(created_parts=0, created_models=0, linked_parts=0, errors=[])

    for i, raw in enumerate(rows, start=2):  # start=2 -> first data row after header
        record = {header[j]: raw[j] for j in range(min(len(header), len(raw)))}

        part_code = _to_str(record.get("part_code"))
        part_name = _to_str(record.get("part_name"))
        if not part_code or not part_name:
            result.errors.append(f"Row {i}: missing part_code/part_name (skipped)")
            continue

        if db.query(Part).filter(Part.part_code == part_code).first():
            result.errors.append(f"Row {i}: part_code '{part_code}' exists (skipped)")
            continue

        # --- resolve / create the linked AI model from its name ---
        ai_model_id = None
        model_name = _to_str(record.get("model_name"))
        if model_name:
            model = db.query(AIModel).filter(AIModel.model_name == model_name).first()
            if model is None:
                model = AIModel(
                    model_name=model_name,
                    model_path=resolve_model_filename(model_name),  # '<name>.pt'
                )
                db.add(model)
                db.flush()  # get model_id
                result.created_models += 1
            ai_model_id = model.model_id

        part = Part(
            part_code=part_code,
            part_name=part_name,
            category_id=_to_int(record.get("category_id")),
            ai_model_id=ai_model_id,
            mode_of_operation=_to_str(record.get("mode_of_operation")) or "Counting",
            parts_metadata=_to_str(record.get("parts_metadata")),
            part_weight=_to_float(record.get("part_weight")),
            part_height=_to_float(record.get("part_height")),
            part_width=_to_float(record.get("part_width")),
            part_inner_diameter=_to_float(record.get("part_inner_diameter")),
            part_outer_diameter=_to_float(record.get("part_outer_diameter")),
            part_length=_to_float(record.get("part_length")),
            part_angle=_to_float(record.get("part_angle")),
            part_arch_length=_to_float(record.get("part_arch_length")),
            part_co_planarity=_to_bool(record.get("part_co_planarity")),
            part_parallelity=_to_bool(record.get("part_parallelity")),
            part_concentricity=_to_bool(record.get("part_concentricity")),
        )
        db.add(part)
        result.created_parts += 1
        if ai_model_id is not None:
            result.linked_parts += 1

    db.commit()
    return result
